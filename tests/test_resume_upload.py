import io

import pytest
from docx import Document

from modules.resume_upload import (
    MAX_RESUME_BYTES,
    build_resume_docx,
    extract_resume_text,
    polish_resume_markdown,
    review_resume,
)


def test_extracts_utf8_text_resume():
    assert extract_resume_text("resume.txt", "Asha Student\nSkills\nPython".encode()) == (
        "Asha Student\nSkills\nPython"
    )


def test_extracts_docx_paragraphs_and_table_content():
    source = Document()
    source.add_paragraph("Asha Student")
    source.add_heading("Education", level=1)
    source.add_paragraph("BCA, Example College")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Project"
    table.cell(0, 1).text = "Learning app"
    content = io.BytesIO()
    source.save(content)

    extracted = extract_resume_text("resume.docx", content.getvalue())

    assert "BCA, Example College" in extracted
    assert "Project | Learning app" in extracted


@pytest.mark.parametrize(
    ("filename", "content", "message"),
    [
        ("resume.rtf", b"resume", "PDF, DOCX, or UTF-8 TXT"),
        ("resume.txt", b"\xff", "UTF-8"),
        ("resume.txt", b"", "empty"),
    ],
)
def test_rejects_unsupported_or_unreadable_resume(filename, content, message):
    with pytest.raises(ValueError, match=message):
        extract_resume_text(filename, content)


def test_rejects_oversized_resume():
    with pytest.raises(ValueError, match="10 MB or smaller"):
        extract_resume_text("resume.txt", b"x" * (MAX_RESUME_BYTES + 1))


def test_review_reports_missing_sections_and_actionable_suggestions():
    review = review_resume("Asha Student\nSkills\nPython, SQL")

    assert review["sections"] == ["Skills"]
    assert not review["has_email"]
    assert any("projects" in suggestion.lower() for suggestion in review["suggestions"])
    assert any(
        "quantify outcomes" in suggestion.lower()
        for suggestion in review["suggestions"]
    )


def test_review_compares_resume_skills_with_selected_target_role():
    review = review_resume(
        "Asha Student\nSkills\nPython, SQL",
        ["Python", "SQL", "Data Visualization"],
    )

    assert review["matched_target_skills"] == ["Python", "SQL"]
    assert review["missing_target_skills"] == ["Data Visualization"]


def test_polished_resume_normalizes_sections_without_dropping_content():
    source = (
        "Asha Student\nasha@example.com\n"
        "WORK EXPERIENCE:\n- Built a learning app for 20 students\n"
        "EDUCATION\nBCA, Example College"
    )

    polished = polish_resume_markdown(source)

    assert polished.startswith("# Asha Student")
    assert "## Experience" in polished
    assert "## Education" in polished
    assert "Built a learning app for 20 students" in polished


def test_builds_downloadable_word_document():
    content = build_resume_docx(
        "# Asha Student\n\nasha@example.com\n\n## Skills\n\n- Python"
    )
    document = Document(io.BytesIO(content))

    assert document.paragraphs[0].text == "Asha Student"
    assert any(paragraph.text == "Skills" for paragraph in document.paragraphs)
    assert any(paragraph.text == "Python" for paragraph in document.paragraphs)
