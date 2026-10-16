"""Resume document extraction, review, and professional formatting."""

from __future__ import annotations

import io
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader

MAX_RESUME_BYTES = 10 * 1024 * 1024

_SECTION_ALIASES = {
    "Professional Summary": {
        "about me", "career objective", "objective", "profile", "professional profile",
        "professional summary", "summary",
    },
    "Experience": {
        "employment", "employment history", "experience", "professional experience",
        "work experience", "work history",
    },
    "Education": {"academic background", "education", "educational background"},
    "Skills": {
        "core competencies", "core skills", "key skills", "skills", "technical skills",
    },
    "Projects": {"academic projects", "key projects", "personal projects", "projects"},
    "Certifications": {"certificates", "certifications", "licenses"},
    "Achievements": {"achievements", "accomplishments", "awards", "honors"},
    "Languages": {"language", "languages"},
}
_HEADING_LOOKUP = {
    re.sub(r"[^a-z0-9 ]", "", alias.lower()).strip(): heading
    for heading, aliases in _SECTION_ALIASES.items()
    for alias in aliases
}
_ACTION_VERBS = {
    "analyzed", "built", "created", "delivered", "designed", "developed",
    "implemented", "improved", "increased", "launched", "led", "managed",
    "optimized", "organized", "reduced", "researched", "resolved", "streamlined",
}


def extract_resume_text(filename: str, content: bytes) -> str:
    """Extract text from a supported resume without persisting its contents."""
    if not isinstance(content, bytes) or not content:
        raise ValueError("The uploaded file is empty.")
    if len(content) > MAX_RESUME_BYTES:
        raise ValueError("Resume files must be 10 MB or smaller.")

    extension = Path(filename).suffix.lower()
    try:
        if extension == ".txt":
            text = content.decode("utf-8-sig")
        elif extension == ".pdf":
            reader = PdfReader(io.BytesIO(content), strict=False)
            if reader.is_encrypted:
                raise ValueError("Password-protected PDFs are not supported.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif extension == ".docx":
            document = Document(io.BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        else:
            raise ValueError("Upload a PDF, DOCX, or UTF-8 TXT resume.")
    except UnicodeDecodeError as error:
        raise ValueError("TXT resumes must use UTF-8 text encoding.") from error
    except ValueError:
        raise
    except Exception as error:
        raise ValueError(
            "The file could not be read. Check that it is a valid, unprotected "
            "PDF, DOCX, or UTF-8 TXT document."
        ) from error

    text = text.replace("\x00", "")
    if not text.strip():
        raise ValueError(
            "No readable text was found. Scanned/image-only PDFs need OCR before upload."
        )
    return text.strip()


def _heading_for_line(line: str) -> str | None:
    candidate = re.sub(r"^[#\s]*|[\s:：]+$", "", line).lower()
    candidate = re.sub(r"[^a-z0-9 ]", "", candidate).strip()
    return _HEADING_LOOKUP.get(candidate)


def _split_resume_sections(text: str) -> tuple[list[str], list[tuple[str, list[str]]]]:
    preamble = []
    sections: list[tuple[str, list[str]]] = []
    current_name = None
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        heading = _heading_for_line(line)
        if heading:
            current_name = heading
            if not any(name == heading for name, _ in sections):
                sections.append((heading, []))
            continue
        if current_name is None:
            preamble.append(line)
        else:
            next(section_lines for name, section_lines in sections if name == current_name).append(line)
    return preamble, sections


def _looks_like_name(line: str) -> bool:
    return (
        len(line) <= 70
        and not any(character in line for character in "@:/|")
        and len(re.findall(r"[A-Za-z][A-Za-z'-]*", line)) in range(2, 6)
        and not re.search(r"\d", line)
    )


def polish_resume_markdown(text: str) -> str:
    """Normalize resume headings and layout without adding or changing facts."""
    preamble, sections = _split_resume_sections(text)
    name = preamble[0] if preamble and _looks_like_name(preamble[0]) else "Professional Resume"
    contact_lines = preamble[1:] if name != "Professional Resume" else preamble

    lines = [f"# {name}"]
    if sections and contact_lines:
        lines.extend(["", " · ".join(contact_lines)])
    elif not sections:
        for item in contact_lines:
            lines.extend(["", item])
    for heading, body in sections:
        lines.extend(["", f"## {heading}"])
        for item in body:
            if item.startswith(("-", "*", "•")):
                lines.append(f"- {item.lstrip('-*• ').strip()}")
            else:
                lines.extend(["", item])
    if not sections and not contact_lines and name == "Professional Resume":
        lines.extend(["", text.strip()])
    return "\n".join(lines).strip() + "\n"


def review_resume(text: str, target_skills: list[str] | None = None) -> dict:
    """Return practical, rule-based resume feedback; does not invent credentials."""
    _, sections = _split_resume_sections(text)
    present = {name for name, body in sections if body}
    normalized = text.lower()
    target_skills = target_skills or []
    matched_target_skills = [
        skill
        for skill in target_skills
        if re.search(rf"(?<!\w){re.escape(skill.lower())}(?!\w)", normalized)
    ]
    missing_target_skills = [
        skill for skill in target_skills if skill not in matched_target_skills
    ]
    bullets = [
        line.strip()
        for line in text.splitlines()
        if re.match(r"^\s*(?:[-*•]|\d+[.)])\s+", line)
    ]
    suggestions = []
    if not re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", text):
        suggestions.append("Add a professional email address in the contact header.")
    if "Experience" not in present:
        suggestions.append(
            "Add an Experience section if you have internships, employment, or relevant volunteering."
        )
    if "Professional Summary" not in present:
        suggestions.append(
            "Add a short, role-focused professional summary based on your real strengths and experience."
        )
    if "Education" not in present:
        suggestions.append("Add an Education section with your degree, institution, and dates.")
    if "Skills" not in present:
        suggestions.append("Add a concise Skills section tailored to the role you want.")
    if "Projects" not in present:
        suggestions.append(
            "Add 2–3 relevant projects and explain your individual contribution and tools used."
        )
    if not bullets:
        suggestions.append(
            "Use short bullet points so recruiters can scan responsibilities and results quickly."
        )
    if not re.search(r"\b\d+(?:[.,]\d+)?\s*(?:%|percent|users|customers|hours|days|weeks|months|years|projects|records|team members)\b|[$₹€]\s?\d+", normalized):
        suggestions.append(
            "Where accurate, quantify outcomes (for example users reached, time saved, or error reduction)."
        )
    if not any(re.search(rf"\b{verb}\w*\b", normalized) for verb in _ACTION_VERBS):
        suggestions.append(
            "Start experience and project bullets with clear action verbs such as Built, Analyzed, or Improved."
        )
    if target_skills and not matched_target_skills:
        suggestions.append(
            "Your resume does not mention the selected role's listed skills. "
            "Include relevant skills only if you can confidently support them."
        )
    if not suggestions:
        suggestions.append(
            "The core sections are present. Proofread dates and tailor keywords to each job description."
        )
    return {
        "word_count": len(re.findall(r"\b[\w'-]+\b", text)),
        "sections": [name for name, body in sections if body],
        "bullet_count": len(bullets),
        "has_email": bool(re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", text)),
        "has_quantified_results": bool(
            re.search(
                r"\b\d+(?:[.,]\d+)?\s*(?:%|percent|users|customers|hours|days|weeks|months|years|projects|records|team members)\b|[$₹€]\s?\d+",
                normalized,
            )
        ),
        "matched_target_skills": matched_target_skills,
        "missing_target_skills": missing_target_skills,
        "suggestions": suggestions,
    }


def build_resume_docx(markdown: str) -> bytes:
    """Build a clean, ATS-readable DOCX using only the supplied resume content."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(45, 55, 65)
    normal.paragraph_format.space_after = Pt(4)

    for line in markdown.splitlines():
        if not line.strip():
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(line[2:].strip())
            run.font.name = "Calibri"
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(28, 55, 90)
        elif line.startswith("## "):
            paragraph = document.add_paragraph(line[3:].strip(), style="Heading 2")
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(4)
        elif line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            document.add_paragraph(line.strip())
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
